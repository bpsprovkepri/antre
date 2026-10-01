"""Uji terhadap PostgreSQL sungguhan (pgserver). Jalankan: pytest -q"""
import socket
import threading
import time
from concurrent.futures import ThreadPoolExecutor

import pytest

from types import SimpleNamespace

PNG = (b"\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR\x00\x00\x00\x01\x00\x00\x00\x01\x08\x06\x00\x00\x00\x1f\x15\xc4\x89"
       b"\x00\x00\x00\rIDATx\x9cc\xf8\xff\xff?\x00\x05\xfe\x02\xfe\xa7\x9a\xa0\xa0\x00\x00\x00\x00IEND\xaeB`\x82")


@pytest.fixture()
def mod(klien):
    """Impor modul aplikasi setelah lingkungan database uji siap."""
    from app import laporan, main, pengaturan, suara, util
    return SimpleNamespace(main=main, pengaturan=pengaturan, suara=suara, util=util, laporan=laporan)


def bersih(k):
    k.srv.psql("DELETE FROM queue_tiket")


# ---------- fungsi murni ----------
def test_terbilang(mod):
    util = mod.util
    assert [util.terbilang(n) for n in (0, 2, 10, 11, 12, 19, 20, 21, 100, 101, 115, 200, 999, 1000, 2025)] == [
        "nol", "dua", "sepuluh", "sebelas", "dua belas", "sembilan belas", "dua puluh", "dua puluh satu", "seratus",
        "seratus satu", "seratus lima belas", "dua ratus", "sembilan ratus sembilan puluh sembilan", "seribu",
        "dua ribu dua puluh lima"]
    assert util.teks_panggil(12, "Meja 2") == "Nomor antrian, dua belas, silakan menuju, Meja dua"


def test_youtube_embed(mod):
    y = mod.pengaturan.youtube_embed
    lama = y("tgdk7wCVLpg?autoplay=1&showinfo=0&loop=1&list=PL88EA9KKUmj1I6LRpec2LlxRw4b7K3DA9&rel=0")
    assert lama.startswith("https://www.youtube-nocookie.com/embed/tgdk7wCVLpg?list=PL88EA9KKUmj1I6LRpec2LlxRw4b7K3DA9&")
    assert "mute=1" in lama and "autoplay=1" in lama and "loop=1" in lama
    assert "playlist=dQw4w9WgXcQ" in y("dQw4w9WgXcQ")
    assert "embed/dQw4w9WgXcQ" in y("https://www.youtube.com/watch?v=dQw4w9WgXcQ&t=5")
    assert "embed/dQw4w9WgXcQ" in y("https://youtu.be/dQw4w9WgXcQ")
    assert "videoseries?list=PL88EA9KKUmj1I6LRpec2LlxRw4b7K3DA9" in y("PL88EA9KKUmj1I6LRpec2LlxRw4b7K3DA9")
    for buruk in ("", "https://evil.example/x", "javascript:alert(1)", "abc def", '"><script>'):
        assert y(buruk) == ""


# ---------- awal ----------
def test_seed_dari_tabel_lama_dan_tabel_lama_utuh(klien, mod):
    pengaturan = mod.pengaturan
    p = pengaturan.get(paksa=True)
    assert p["nama_instansi"] == "Pelayanan Statistik Terpadu BPS Provinsi Kepulauan Riau"
    assert p["alamat"].startswith("Jl. Ahmad Yani") and p["telpon"] == "07714500155"
    assert [m["no"] for m in p["meja"]] == [1, 2] and p["meja"][1]["nama"] == "Meja 2"
    assert p["warna_accent"] == "#ff00e6" and p["warna_background"] == "#212529"
    assert pengaturan.youtube_embed(p["youtube_id"]).startswith("https://www.youtube-nocookie.com/embed/tgdk7wCVLpg")
    assert "queue_tiket" in klien.srv.psql("select tablename from pg_tables where schemaname='public'")
    assert klien.srv.psql("select count(*) from queue_antrian_admisi").split()[2] == "1"   # tabel lama tidak disentuh
    assert klien.get("/healthz").json() == {"status": "ok"}


def test_halaman_publik(klien):
    r = klien.get("/")
    assert r.status_code == 200
    for teks in ("Nomor Antrian", "Panggilan Antrian", "Monitor Antrian", "Setting Antrian", "Laporan Layanan", "Kelola Data", "Meja 2"):
        assert teks in r.text
    k = klien.get("/kiosk")
    assert k.status_code == 200 and "Ambil Nomor" in k.text and 'href="/"' in k.text     # ada navigasi ke home
    badan = k.text.split("<script>")[0]                                         # bagian yang terlihat pengunjung
    assert "bi-printer" not in badan and "Cetak" not in badan and "print(" not in badan  # tombol cetak disembunyikan
    m = klien.get("/monitor")
    assert m.status_code == 200
    for teks in ("Nomor Antrian Sekarang", "Antrian Selanjutnya", "Sisa Antrian", "youtube-nocookie.com/embed/tgdk7wCVLpg", "SELAMAT DATANG DI PST", 'href="/"'):
        assert teks in m.text
    assert "Total Antrian" not in m.text
    assert klien.get("/static/vendor/bootstrap/bootstrap.min.css").status_code == 200
    assert klien.get("/logo").headers["content-type"].startswith("image/svg")


# ---------- login ----------
def test_login_sesi(klien):
    klien.cookies.clear()
    r = klien.get("/loket/1", follow_redirects=False)
    assert r.status_code == 303 and r.headers["location"].startswith("/login?next=/loket/1")
    assert klien.get("/pengaturan", follow_redirects=False).status_code == 303
    assert klien.get("/laporan", follow_redirects=False).status_code == 303
    assert klien.get("/api/loket/1").status_code == 401
    assert klien.post("/api/pengaturan", data={}).status_code == 401
    bad = klien.post("/login", data={"username": "admin", "password": "salah", "next": "/"})
    assert bad.status_code == 401 and "salah" in bad.text
    ok = klien.post("/login", data={"username": "admin", "password": "rahasia-test", "next": "//evil.com"}, follow_redirects=False)
    assert ok.status_code == 303 and ok.headers["location"] == "/"          # tidak bisa dialihkan ke situs lain
    assert klien.get("/pengaturan").status_code == 200
    assert klien.post("/logout", follow_redirects=False).status_code == 303
    assert klien.get("/pengaturan", follow_redirects=False).status_code == 303


def test_batas_percobaan_login(klien, mod):
    klien.cookies.clear()
    kode = [klien.post("/login", data={"username": "admin", "password": "x%d" % i, "next": "/"}).status_code for i in range(7)]
    assert 429 in kode
    mod.main._gagal.clear()


# ---------- alur antrian ----------
def test_alur_panggil_mulai_selesai(admin):
    k = admin
    bersih(k)
    nomor = [k.post("/api/tiket").json()["nomor"] for _ in range(3)]
    assert nomor == [1, 2, 3]
    s = k.get("/api/loket/1").json()
    rows = {r["nomor"]: r for r in s["rows"]}
    assert s["jumlah"] == 3 and s["sisa"] == 3 and s["selanjutnya"] == "001"
    assert rows["001"]["bisa_panggil"] and not rows["002"]["bisa_panggil"] and not rows["003"]["bisa_panggil"]
    ids = {n: r["id"] for n, r in rows.items()}
    do = lambda m, n, a: k.post(f"/api/loket/{m}/{ids[n]}/{a}")
    assert do(1, "002", "panggil").status_code == 409            # tidak boleh melompat nomor
    assert do(1, "001", "selesai").status_code == 409            # belum dipanggil
    assert do(1, "001", "mulai").status_code == 404              # aksi 'mulai' sudah digabung ke 'panggil'
    r = do(1, "001", "panggil"); assert r.status_code == 200
    j = r.json(); assert j["teks"] == "Nomor antrian, satu, silakan menuju, Meja satu" and j["n"] == 1
    row = k.srv.psql("select dipanggil_at is not null, mulai_at is not null, dipanggil_at = mulai_at from queue_tiket where nomor=1")
    assert "t|t|t" in "".join(row.split())                                   # satu klik = panggil DAN mulai layanan
    s = k.get("/api/loket/1").json(); rows = {r["nomor"]: r for r in s["rows"]}
    assert s["saya"] == "001" and rows["001"]["status"] == "aktif" and rows["001"]["bisa_selesai"] and rows["001"]["bisa_ulang"]
    assert not rows["002"]["bisa_panggil"]                        # 002 nonaktif selama 001 belum selesai
    assert s["rows"][0]["nomor"] == "001"                         # yang aktif tampil paling atas
    assert do(1, "002", "panggil").status_code == 409
    assert do(1, "001", "panggil").json()["n"] == 2               # panggil ulang tidak mengubah waktu mulai
    assert do(2, "001", "selesai").status_code == 409             # meja lain tidak bisa mengambil alih
    assert do(2, "002", "panggil").status_code == 200             # meja 2 bebas -> ambil 002
    time.sleep(1.1)
    assert do(1, "001", "selesai").status_code == 200
    assert do(1, "001", "selesai").status_code == 409
    s = k.get("/api/loket/1").json(); rows = {r["nomor"]: r for r in s["rows"]}
    assert rows["001"]["status"] == "selesai" and rows["001"]["durasi"] >= "00:01"
    assert rows["003"]["bisa_panggil"]                            # kini 003 boleh dipanggil meja 1
    assert do(1, "003", "panggil").status_code == 200
    assert do(1, "003", "lewati").status_code == 200              # tidak hadir
    s = k.get("/api/loket/1").json(); rows = {r["nomor"]: r for r in s["rows"]}
    assert rows["003"]["status"] == "lewat"
    assert k.post("/api/loket/9/1/panggil").status_code == 404
    assert k.post(f"/api/loket/1/{ids['001']}/hapus").status_code == 404


def test_monitor_api(klien):
    d = klien.get("/api/monitor").json()          # public, tanpa login
    assert d["last"]["nomor_txt"] == "003" and d["last"]["meja_nama"] == "Meja 1"
    assert d["last"]["teks"] == "Nomor antrian, tiga, silakan menuju, Meja satu"
    assert d["sisa"] == 0 and d["selanjutnya"] == "-"
    assert [a["nomor"] for a in d["aktif"]] == ["002"]      # 002 masih berjalan di Meja 2
    assert d["suara"] == {"aktif": True, "engine": "auto"} and d["versi"]


RIWAYAT = """
DELETE FROM queue_tiket;
INSERT INTO queue_tiket(tanggal, nomor, meja, dibuat_at, dipanggil_at, mulai_at, selesai_at, lewat) VALUES
 ('2025-12-30',1,1,'2025-12-30 09:00:00+07','2025-12-30 09:01:00+07','2025-12-30 09:01:00+07','2025-12-30 09:07:00+07',false),
 ('2026-08-03',1,1,'2026-08-03 10:00:00+07','2026-08-03 10:02:00+07','2026-08-03 10:02:00+07','2026-08-03 10:07:00+07',false),
 ('2026-08-03',2,2,'2026-08-03 10:05:00+07','2026-08-03 10:10:00+07','2026-08-03 10:10:00+07','2026-08-03 10:12:00+07',false),
 ('2026-08-03',3,1,'2026-08-03 10:20:00+07','2026-08-03 10:30:00+07','2026-08-03 10:30:00+07','2026-08-03 10:31:00+07',true),
 ('2026-08-04',1,1,'2026-08-04 08:00:00+07','2026-08-04 08:01:00+07','2026-08-04 08:01:00+07','2026-08-04 08:04:00+07',false),
 ('2026-08-04',2,NULL,'2026-08-04 08:30:00+07',NULL,NULL,NULL,false),
 ('2026-09-15',1,2,'2026-09-15 13:00:00+07','2026-09-15 13:01:00+07','2026-09-15 13:01:00+07','2026-09-15 13:02:00+07',false);
"""


def hitung(mod, **kw):
    return mod.laporan.data(mod.laporan.periode(**kw), mod.pengaturan.get())


@pytest.fixture()
def riwayat(klien):
    klien.srv.psql(RIWAYAT)
    return klien


def test_laporan_wajib_login_semua_jenis(riwayat):
    k = riwayat
    k.cookies.clear()
    for url in ("/laporan", "/laporan?jenis=harian&tanggal=2026-08-03", "/laporan?jenis=bulanan&bulan=8&tahun=2026",
                "/laporan?jenis=tahunan&tahun=2026", "/laporan?jenis=rentang&awal=2026-08-01&akhir=2026-08-31",
                "/laporan.xlsx", "/laporan.xlsx?jenis=bulanan&bulan=8&tahun=2026", "/laporan.xlsx?jenis=tahunan&tahun=2026",
                "/laporan.xlsx?jenis=rentang&awal=2026-08-01&akhir=2026-08-31"):
        r = k.get(url, follow_redirects=False)
        assert r.status_code == 303 and r.headers["location"].startswith("/login"), url
        assert "Agustus" not in r.text and "Total antrian" not in r.text       # tidak ada data yang bocor


def test_hitungan_harian_bulanan_tahunan_rentang(riwayat, mod):
    h = hitung(mod, jenis="harian", tanggal=__import__("datetime").date(2026, 8, 3))["ring"]
    assert (h["total"], h["selesai"], h["lewat"], h["rata_layanan"], h["maks_layanan"], h["min_layanan"]) == (3, 2, 1, "03:30", "05:00", "02:00")
    b = hitung(mod, jenis="bulanan", bulan=8, tahun=2026)
    r = b["ring"]
    assert (r["total"], r["selesai"], r["lewat"], r["tdk_dilayani"], r["hari_aktif"], r["rata_per_hari"]) == (5, 3, 1, 1, 2, 2.5)
    assert (r["rata_layanan"], r["maks_layanan"], r["min_layanan"], r["rata_tunggu"]) == ("03:20", "05:00", "02:00", "04:30")
    assert b["gran"] == "hari" and len(b["seri"]) == 31 and b["seri"][2]["n"] == 3 and b["seri"][3]["n"] == 2
    assert b["seri"][2]["rata"] == "03:30" and b["seri"][2]["label_panjang"] == "Sen, 3 Agustus 2026"
    t = hitung(mod, jenis="tahunan", tahun=2026)
    assert t["ring"]["total"] == 6 and t["ring"]["selesai"] == 4 and t["ring"]["rata_layanan"] == "02:45"
    assert t["gran"] == "bulan" and len(t["seri"]) == 12 and [s["n"] for s in t["seri"]] == [0, 0, 0, 0, 0, 0, 0, 5, 1, 0, 0, 0]
    assert hitung(mod, jenis="tahunan", tahun=2025)["ring"]["total"] == 1                       # Desember 2025 tidak ikut 2026
    assert hitung(mod, jenis="tahunan", tahun=2024)["ring"]["total"] == 0                       # periode kosong tetap aman
    D = __import__("datetime").date
    a = hitung(mod, jenis="rentang", awal=D(2026, 8, 3), akhir=D(2026, 8, 4))["ring"]
    assert a["total"] == 5 and a["rata_layanan"] == "03:20"
    assert hitung(mod, jenis="rentang", awal=D(2026, 8, 4), akhir=D(2026, 8, 4))["ring"]["total"] == 2
    assert hitung(mod, jenis="rentang", awal=D(2026, 8, 5), akhir=D(2026, 9, 14))["ring"]["total"] == 0   # batas tanggal inklusif
    lebar = hitung(mod, jenis="rentang", awal=D(2026, 7, 1), akhir=D(2026, 10, 31))
    assert lebar["gran"] == "bulan" and len(lebar["seri"]) == 4 and lebar["ring"]["total"] == 6          # >62 hari dikelompokkan per bulan
    assert lebar["seri"][1]["label"] == "Agu"


def test_periode_dan_navigasi(mod):
    lap, D = mod.laporan, __import__("datetime").date
    for kw, pesan in (({"jenis": "bulanan", "bulan": 13, "tahun": 2026}, "Bulan"), ({"jenis": "bulanan", "bulan": 1, "tahun": 1999}, "Tahun"),
                      ({"jenis": "tahunan", "tahun": 2500}, "Tahun"),
                      ({"jenis": "rentang", "awal": D(2026, 8, 5), "akhir": D(2026, 8, 1)}, "lebih awal"),
                      ({"jenis": "rentang", "awal": D(2025, 1, 1), "akhir": D(2026, 8, 1)}, "maksimal")):
        with pytest.raises(lap.PeriodeError, match=pesan):
            lap.periode(**kw)
    assert lap.periode("ngawur")["jenis"] == "harian"
    per = lap.periode("rentang", awal=D(2026, 1, 1), akhir=D(2026, 12, 31))        # 365 hari masih boleh
    assert per["hari"] == 365
    n = lap.navigasi(lap.periode("bulanan", bulan=1, tahun=2026))
    assert "bulan=12" in n["prev"] and "tahun=2025" in n["prev"] and "bulan=2" in n["next"]
    assert "bulan=1" in lap.navigasi(lap.periode("bulanan", bulan=12, tahun=2025))["next"]
    assert lap.navigasi(lap.periode("harian", tanggal=D(2026, 8, 3)))["prev"].endswith("2026-08-02")
    assert lap.navigasi(lap.periode("harian"))["next"] is None                      # tidak bisa maju ke masa depan
    assert lap.navigasi(lap.periode("tahunan", tahun=2025))["prev"].endswith("2024")
    assert len(lap.navigasi(lap.periode("rentang"))["preset"]) == 4


def test_halaman_laporan_semua_jenis(riwayat):
    k = riwayat
    r = k.post("/login", data={"username": "admin", "password": "rahasia-test", "next": "/"}, follow_redirects=False)
    assert r.status_code == 303
    h = k.get("/laporan?jenis=harian&tanggal=2026-08-03")
    assert h.status_code == 200 and "Laporan Harian" in h.text and "Senin, 3 Agustus 2026" in h.text and "03:30" in h.text
    for teks in ("Harian", "Bulanan", "Tahunan", "Rentang Waktu", "Unduh Excel", 'href="/"', 'action="/logout"', "Detail antrian", "Tidak hadir"):
        assert teks in h.text
    b = k.get("/laporan?jenis=bulanan&bulan=8&tahun=2026")
    assert b.status_code == 200 and "Agustus 2026" in b.text and 'name="bulan"' in b.text and 'name="tahun"' in b.text
    assert '<option value="8" selected>Agustus</option>' in b.text and '<option value="2025"' in b.text
    assert "Rincian per hari" in b.text and "laporan.xlsx?jenis=bulanan&amp;bulan=8&amp;tahun=2026" in b.text.replace("&", "&amp;").replace("&amp;amp;", "&amp;")
    t = k.get("/laporan?jenis=tahunan&tahun=2026")
    assert t.status_code == 200 and "Tahun 2026" in t.text and "Rincian per bulan" in t.text and "Agustus 2026" in t.text
    g = k.get("/laporan?jenis=rentang&awal=2026-08-03&akhir=2026-08-04")
    assert g.status_code == 200 and 'name="awal"' in g.text and 'name="akhir"' in g.text and "Pintasan" in g.text and "7 hari terakhir" in g.text
    # input tidak valid -> pesan jelas, halaman tetap tampil (tidak error 500)
    for url, pesan in (("/laporan?jenis=rentang&awal=2026-08-05&akhir=2026-08-01", "tidak boleh lebih awal"),
                       ("/laporan?jenis=rentang&awal=2020-01-01&akhir=2026-08-01", "maksimal 366"),
                       ("/laporan?jenis=bulanan&bulan=13&tahun=2026", "Bulan harus"), ("/laporan?jenis=tahunan&tahun=1999", "Tahun harus")):
        x = k.get(url)
        assert x.status_code == 422 and pesan in x.text and "Laporan Harian" in x.text, url
    assert k.get("/laporan?jenis=ngawur").status_code == 200
    assert k.get("/laporan?jenis=harian&tanggal=bukan-tanggal").status_code == 422
    assert k.get("/laporan?jenis=bulanan&bulan=abc").status_code == 422


def test_laporan_detail_banyak_hanya_di_excel(riwayat, mod, monkeypatch):
    monkeypatch.setattr(mod.laporan, "MAKS_DETAIL_LAYAR", 2)
    d = hitung(mod, jenis="bulanan", bulan=8, tahun=2026)
    assert d["tampil_detail"] is False
    r = riwayat.get("/laporan?jenis=bulanan&bulan=8&tahun=2026")
    assert "terlalu banyak untuk ditampilkan" in r.text and "lembar <b>Detail</b>" in r.text


def test_excel_semua_jenis(riwayat):
    import io
    from openpyxl import load_workbook
    k = riwayat
    x = k.get("/laporan.xlsx?jenis=bulanan&bulan=8&tahun=2026")
    assert x.status_code == 200 and x.content[:2] == b"PK" and "spreadsheetml" in x.headers["content-type"]
    assert 'filename="laporan-antrian-bulanan-2026-08.xlsx"' in x.headers["content-disposition"]
    wb = load_workbook(io.BytesIO(x.content))
    assert wb.sheetnames == ["Ringkasan", "Rincian", "Detail"]
    ring = {r[0]: r[1] for r in wb["Ringkasan"].iter_rows(values_only=True) if r[0]}
    assert ring["Laporan Bulanan"] == "Agustus 2026" and ring["Total antrian"] == 5 and ring["Rata-rata lama layanan (mm:ss)"] == "03:20"
    assert wb["Rincian"].max_row == 32                                   # header + 31 hari
    det = list(wb["Detail"].iter_rows(values_only=True))
    assert len(det) == 6 and det[1][1] == "001" and det[1][9] == 300 and det[3][10] == "tidak hadir" and det[5][10] == "tidak dilayani"
    assert det[1][0].strftime("%Y-%m-%d") == "2026-08-03"
    t = load_workbook(io.BytesIO(k.get("/laporan.xlsx?jenis=tahunan&tahun=2026").content))
    assert t["Rincian"].max_row == 13 and t["Detail"].max_row == 7
    g = k.get("/laporan.xlsx?jenis=rentang&awal=2026-08-03&akhir=2026-08-03")
    assert 'laporan-antrian-rentang-2026-08-03_2026-08-03.xlsx' in g.headers["content-disposition"]
    assert load_workbook(io.BytesIO(g.content))["Detail"].max_row == 4
    h = load_workbook(io.BytesIO(k.get("/laporan.xlsx?jenis=harian&tanggal=2026-08-03").content))
    assert h["Rincian"].max_row == 11                                     # per jam 07-16
    assert k.get("/laporan.xlsx?jenis=rentang&awal=2026-08-05&akhir=2026-08-01").status_code == 422


def test_laporan_hari_ini_status_berjalan(admin):
    k = admin
    bersih(k)
    for _ in range(2):
        k.post("/api/tiket")
    rows = {r["nomor"]: r for r in k.get("/api/loket/1").json()["rows"]}
    k.post(f"/api/loket/1/{rows['001']['id']}/panggil")
    page = k.get("/laporan")
    assert page.status_code == 200 and "Berjalan" in page.text and "Menunggu" in page.text and "Hari ini" in page.text
    assert k.get("/laporan.xlsx").status_code == 200


def test_nomor_unik_paralel(klien):
    bersih(klien)
    with ThreadPoolExecutor(10) as ex:
        hasil = list(ex.map(lambda _: klien.post("/api/tiket").json()["nomor"], range(25)))
    assert sorted(hasil) == list(range(1, 26))


# ---------- pengaturan ----------
def form(admin_klien, **ubah):
    from app import pengaturan
    p = pengaturan.get(paksa=True)
    d = {"nama_instansi": p["nama_instansi"], "alamat": p["alamat"], "telpon": p["telpon"], "email": p["email"],
         "running_text": p["running_text"], "youtube_id": p["youtube_id"], "judul_monitor": p["judul_monitor"],
         "footer_text": p["footer_text"], "warna_primary": p["warna_primary"], "warna_secondary": p["warna_secondary"],
         "warna_accent": p["warna_accent"], "warna_background": p["warna_background"], "warna_text": p["warna_text"],
         "suara_loket": "1", "suara_monitor": "1", "suara_engine": "auto", "printer_host": "", "printer_port": "9100",
         "printer_lebar": "48", "printer_footer": "TERIMA KASIH",
         "meja_no": [str(m["no"]) for m in p["meja"]], "meja_nama": [m["nama"] for m in p["meja"]]}
    d.update(ubah)
    return d


def test_halaman_pengaturan(admin):
    r = admin.get("/pengaturan")
    assert r.status_code == 200
    for teks in ("Informasi Instansi", "Daftar Meja", "Styling Monitor", "Suara Panggilan", "Printer Nomor Antrian", "Footer", "Logout", "Simpan", 'href="/"'):
        assert teks in r.text
    assert admin.get("/loket/1").status_code == 200 and "Panggil &amp; Mulai" in admin.get("/loket/1").text
    assert admin.get("/loket/9").status_code == 404


def test_simpan_footer_kustom_dan_escape(admin, mod):
    util = mod.util
    footer = "Made with ❤️ Tim Teknologi Informasi BPS Provinsi Kepulauan Riau"
    versi0 = admin.get("/api/monitor").json()["versi"]
    r = admin.post("/api/pengaturan", data=form(admin, footer_text=footer, judul_monitor="Monitor <b>PST</b>"))
    assert r.status_code == 200 and r.json()["ok"]
    home = admin.get("/").text
    assert footer in home                                          # footer beranda bisa diubah dari pengaturan
    assert footer in admin.get("/monitor").text
    assert "Monitor &lt;b&gt;PST&lt;/b&gt;" in admin.get("/monitor").text      # HTML dari pengaturan di-escape
    assert admin.get("/api/monitor").json()["versi"] != versi0                 # layar monitor akan memuat ulang
    admin.post("/api/pengaturan", data=form(admin, footer_text="© {tahun} {instansi} <script>x</script>"))
    home = admin.get("/").text
    assert "<script>x</script>" not in home and "&lt;script&gt;x&lt;/script&gt;" in home and str(util.sekarang().year) in home


def test_validasi_pengaturan(admin, mod):
    pengaturan = mod.pengaturan
    def salah(**u):
        r = admin.post("/api/pengaturan", data=form(admin, **u))
        assert r.status_code == 422 and r.json()["errors"], u
        return " ".join(r.json()["errors"])
    assert "warna" in salah(warna_primary="merah")
    assert "YouTube" in salah(youtube_id="https://evil.example/x")
    assert "Nama instansi" in salah(nama_instansi=" ")
    assert "Minimal" in salah(meja_no=[], meja_nama=[])
    assert "dobel" in salah(meja_no=["1", "1"], meja_nama=["a", "b"])
    assert "Port" in salah(printer_port="99999")
    assert "Alamat printer" in salah(printer_aktif="1", printer_host="")
    assert "Alamat printer" in salah(printer_host="a b;rm -rf")
    assert pengaturan.get(paksa=True)["warna_primary"] == "#0a6b38"   # penolakan tidak menyimpan apa pun


def test_meja_dinamis(admin):
    r = admin.post("/api/pengaturan", data=form(admin, meja_no=["1", "2", "3"], meja_nama=["Meja 1", "Meja 2", "Loket VIP 3"]))
    assert r.status_code == 200
    assert "Loket VIP 3" in admin.get("/").text and admin.get("/loket/3").status_code == 200
    admin.post("/api/pengaturan", data=form(admin, meja_no=["1", "2"], meja_nama=["Meja 1", "Meja 2"]))
    assert admin.get("/loket/3").status_code == 404


def test_logo(admin):
    files = lambda b: {"logo": ("logo.png", b, "image/png")}
    r = admin.post("/api/pengaturan", data=form(admin), files=files(b"bukan gambar"))
    assert r.status_code == 422 and "Logo harus" in r.json()["errors"][0]
    r = admin.post("/api/pengaturan", data=form(admin), files=files(b"\x89PNG\r\n\x1a\n" + b"0" * (2 * 1024 * 1024)))
    assert r.status_code == 422 and "2 MB" in r.json()["errors"][0]
    assert admin.post("/api/pengaturan", data=form(admin), files=files(PNG)).status_code == 200
    g = admin.get("/logo")
    assert g.headers["content-type"] == "image/png" and g.content == PNG
    assert "/logo?v=" in admin.get("/monitor").text
    assert admin.post("/api/pengaturan/logo/hapus").status_code == 200
    assert admin.get("/logo").headers["content-type"].startswith("image/svg")


# ---------- printer ----------
class PrinterPalsu:
    def __init__(self):
        self.s = socket.socket(); self.s.bind(("127.0.0.1", 0)); self.s.listen(5)
        self.port = self.s.getsockname()[1]; self.data = []
        threading.Thread(target=self._loop, daemon=True).start()

    def _loop(self):
        while True:
            try:
                c, _ = self.s.accept()
            except OSError:
                return
            buf = b""
            while (x := c.recv(4096)):
                buf += x
            c.close(); self.data.append(buf)


def test_printer_nonaktif_bawaan_dan_cetak_saat_aktif(admin, mod):
    pengaturan = mod.pengaturan
    pr = PrinterPalsu()
    bersih(admin)
    t = admin.post("/api/tiket").json()
    assert t["cetak"] is False and pr.data == []                          # nonaktif: tidak ada yang dicetak

    r = admin.post("/api/pengaturan/tes-cetak", data=form(admin, printer_host="127.0.0.1", printer_port=str(pr.port)))
    assert r.status_code == 200 and r.json()["ok"]
    time.sleep(0.3); assert len(pr.data) == 1 and b"\x1d\x56" in pr.data[0]

    assert admin.post("/api/pengaturan", data=form(admin, printer_aktif="1", printer_host="127.0.0.1",
                                                    printer_port=str(pr.port), printer_footer="TERIMA KASIH")).status_code == 200
    t = admin.post("/api/tiket").json()
    assert t["cetak"] is True
    time.sleep(0.4)
    tiket = pr.data[-1]
    assert b"002" in tiket and b"NOMOR ANTRIAN ANDA" in tiket and b"TERIMA KASIH" in tiket and tiket.endswith(b"\x1d\x56\x42\x00")
    assert "Pelayanan Statistik Terpadu".encode() in tiket

    # printer mati -> nomor tetap terbit, tidak error
    admin.post("/api/pengaturan", data=form(admin, printer_aktif="1", printer_host="127.0.0.1", printer_port="1"))
    assert admin.post("/api/tiket").status_code == 200
    r = admin.post("/api/pengaturan/tes-cetak", data=form(admin, printer_host="127.0.0.1", printer_port="1"))
    assert r.status_code == 502 and "Gagal" in r.json()["pesan"]
    admin.post("/api/pengaturan", data=form(admin))                         # kembalikan: printer nonaktif
    assert pengaturan.get(paksa=True)["printer_aktif"] is False


# ---------- suara ----------
def test_suara_server(klien, mod):
    if not mod.suara.tersedia():
        pytest.skip("espeak-ng tidak terpasang")
    r = klien.get("/api/suara?nomor=12&meja=2")
    assert r.status_code == 200 and r.headers["content-type"] == "audio/wav" and r.content[:4] == b"RIFF" and len(r.content) > 5000
    assert klien.get("/api/suara?nomor=0&meja=1").status_code == 400
    assert klien.get("/api/suara?nomor=1&meja=abc").status_code == 422


def test_halaman_memuat_pemutar_suara(admin):
    assert "Suara.buka()" in admin.get("/loket/1").text and "Suara.panggil" in admin.get("/loket/1").text
    m = admin.get("/monitor").text
    assert "Aktifkan Suara" in m and "Suara.panggil" in m
    assert admin.get("/static/js/suara.js").status_code == 200


def test_sintaks_javascript(admin):
    """Semua skrip inline & berkas .js harus lolos pemeriksaan sintaks Node (dilewati bila node tidak ada)."""
    import re
    import shutil
    import subprocess
    import tempfile
    if not shutil.which("node"):
        pytest.skip("node tidak terpasang")
    daftar = []
    for url in ("/", "/kiosk", "/monitor", "/loket/1", "/laporan", "/pengaturan", "/data"):
        for i, js in enumerate(re.findall(r"<script(?![^>]*src)[^>]*>(.*?)</script>", admin.get(url).text, re.S)):
            daftar.append((f"{url}#{i}", js))
    for f in ("app.js", "suara.js"):
        daftar.append((f, admin.get(f"/static/js/{f}").text))
    assert len(daftar) >= 7
    for nama, js in daftar:
        with tempfile.NamedTemporaryFile("w", suffix=".js", delete=False) as t:
            t.write(js)
        r = subprocess.run(["node", "--check", t.name], capture_output=True, text=True)
        assert r.returncode == 0, f"{nama}: {r.stderr}"


# ---------- kelola data (edit & hapus) ----------
def _seed_kemarin(k):
    """3 tiket kemarin: 001 selesai (meja 1), 002 selesai (meja 2), 003 menunggu. Mengembalikan (tanggal, {nomor: id})."""
    from datetime import timedelta

    from app.util import hari_ini
    t = hari_ini() - timedelta(days=1)
    k.srv.psql(f"""
      DELETE FROM queue_tiket;
      INSERT INTO queue_tiket(tanggal, nomor, meja, dibuat_at, dipanggil_at, mulai_at, selesai_at, lewat, panggil_n, last_call_at) VALUES
       ('{t}',1,1,'{t} 09:00:00+07','{t} 09:02:00+07','{t} 09:02:00+07','{t} 09:07:00+07',false,1,'{t} 09:02:00+07'),
       ('{t}',2,2,'{t} 09:05:00+07','{t} 09:10:00+07','{t} 09:10:00+07','{t} 09:12:00+07',false,1,'{t} 09:10:00+07');
      INSERT INTO queue_tiket(tanggal, nomor, dibuat_at) VALUES ('{t}',3,'{t} 09:20:00+07');
    """)
    rows = k.get(f"/api/data?tanggal={t}").json()["rows"]
    return t, {r["nomor"]: r["id"] for r in rows}


def _ubah(k, id_, **kw):
    return k.post(f"/api/data/{id_}/ubah", json=kw)


def test_kelola_wajib_login(klien):
    klien.cookies.clear()
    r = klien.get("/data", follow_redirects=False)
    assert r.status_code == 303 and r.headers["location"].startswith("/login?next=/data")
    assert klien.get("/api/data").status_code == 401
    assert klien.post("/api/data/1/ubah", json={"status": "menunggu"}).status_code == 401
    assert klien.post("/api/data/hapus", json={"ids": [1]}).status_code == 401
    assert klien.post("/api/data/hapus-tanggal", json={"tanggal": "2026-01-01"}).status_code == 401


def test_halaman_kelola_data(admin):
    r = admin.get("/data")
    assert r.status_code == 200
    for teks in ("Kelola Data Antrian", "Hapus terpilih", "Hapus semua tanggal ini", "Simpan perubahan", 'href="/"', "Meja 2"):
        assert teks in r.text
    assert admin.get("/data?tanggal=2026-08-03").status_code == 200
    assert admin.get("/data?tanggal=bukan-tanggal").status_code == 422
    assert "Kelola Data" in admin.get("/").text and 'href="/data"' in admin.get("/").text


def test_kelola_ubah_waktu_dan_status(admin):
    from app import db
    k = admin
    t, ids = _seed_kemarin(k)
    rows = {r["nomor"]: r for r in k.get(f"/api/data?tanggal={t}").json()["rows"]}
    assert [rows[n]["status"] for n in (1, 2, 3)] == ["selesai", "selesai", "menunggu"]
    assert rows[1]["durasi"] == "05:00" and rows[1]["dipanggil_iso"] == f"{t}T09:02:00" and rows[3]["meja_nama"] == "-"

    # koreksi waktu: lama layanan berubah, mulai ikut waktu dipanggil
    r = _ubah(k, ids[1], status="selesai", meja=1, dibuat=f"{t}T09:00:00", dipanggil=f"{t}T09:03:00", selesai=f"{t}T09:09:30")
    assert r.status_code == 200 and r.json()["nomor_txt"] == "001"
    d = k.get(f"/api/data?tanggal={t}").json()["rows"][0]
    assert d["durasi"] == "06:30" and d["dipanggil"] == "09:03:00" and d["selesai"] == "09:09:30"
    row = db.q("SELECT dipanggil_at = mulai_at AS sama, lewat FROM queue_tiket WHERE id=%s", (ids[1],), one=True)
    assert row["sama"] is True and row["lewat"] is False

    # selesai -> tidak hadir -> menunggu (kolom ikut dibersihkan)
    assert _ubah(k, ids[2], status="lewat", meja=2, dibuat=f"{t}T09:05:00", dipanggil=f"{t}T09:10:00", selesai=f"{t}T09:11:00").status_code == 200
    assert {r["nomor"]: r["status"] for r in k.get(f"/api/data?tanggal={t}").json()["rows"]}[2] == "lewat"
    assert _ubah(k, ids[2], status="menunggu", meja=2, dibuat=f"{t}T09:05:00", dipanggil=f"{t}T09:10:00", selesai=f"{t}T09:11:00").status_code == 200
    row = db.q("SELECT meja, dipanggil_at, mulai_at, selesai_at, lewat, panggil_n, last_call_at FROM queue_tiket WHERE id=%s", (ids[2],), one=True)
    assert row == {"meja": None, "dipanggil_at": None, "mulai_at": None, "selesai_at": None, "lewat": False, "panggil_n": 0, "last_call_at": None}

    # menunggu -> selesai (data yang tadinya belum dilayani dilengkapi)
    assert _ubah(k, ids[3], status="selesai", meja=2, dibuat=f"{t}T09:20:00", dipanggil=f"{t}T09:25:00", selesai=f"{t}T09:28:00").status_code == 200
    row = db.q("SELECT meja, panggil_n, last_call_at IS NOT NULL AS ada FROM queue_tiket WHERE id=%s", (ids[3],), one=True)
    assert row == {"meja": 2, "panggil_n": 1, "ada": True}


def test_kelola_ubah_validasi(admin):
    from datetime import timedelta

    from app.util import hari_ini
    k = admin
    t, ids = _seed_kemarin(k)
    dasar = dict(status="selesai", meja=1, dibuat=f"{t}T09:00:00", dipanggil=f"{t}T09:02:00", selesai=f"{t}T09:07:00")
    tolak = lambda **kw: _ubah(k, ids[1], **{**dasar, **kw})
    r = tolak(selesai=f"{t}T09:01:00"); assert r.status_code == 422 and "selesai tidak boleh lebih awal" in r.json()["detail"]
    r = tolak(dipanggil=f"{t}T08:00:00"); assert r.status_code == 422 and "dipanggil tidak boleh lebih awal" in r.json()["detail"]
    r = tolak(selesai=""); assert r.status_code == 422 and "selesai wajib" in r.json()["detail"]
    r = tolak(dibuat=""); assert r.status_code == 422 and "ambil nomor wajib" in r.json()["detail"]
    r = tolak(dibuat="bukan waktu"); assert r.status_code == 422 and "tidak valid" in r.json()["detail"]
    r = tolak(meja=99); assert r.status_code == 422 and "Meja tidak ditemukan" in r.json()["detail"]
    r = tolak(meja=None); assert r.status_code == 422 and "meja" in r.json()["detail"].lower()
    r = tolak(status="ngawur"); assert r.status_code == 422 and "Status tidak dikenal" in r.json()["detail"]
    besok = hari_ini() + timedelta(days=1)
    r = tolak(selesai=f"{besok}T09:07:00"); assert r.status_code == 422 and ("masa depan" in r.json()["detail"] or "tanggal tiket" in r.json()["detail"])
    r = tolak(selesai=f"{t + timedelta(days=1)}T00:10:00"); assert r.status_code == 422 and "tanggal tiket" in r.json()["detail"]
    assert _ubah(k, 999999999, **dasar).status_code == 404
    # tidak ada yang berubah akibat penolakan
    d = k.get(f"/api/data?tanggal={t}").json()["rows"][0]
    assert d["dipanggil"] == "09:02:00" and d["selesai"] == "09:07:00"

    # satu meja tidak boleh melayani dua nomor sekaligus
    assert _ubah(k, ids[1], **{**dasar, "status": "aktif"}).status_code == 200            # 001 sedang dilayani di meja 1
    r = _ubah(k, ids[3], status="aktif", meja=1, dibuat=f"{t}T09:20:00", dipanggil=f"{t}T09:25:00")
    assert r.status_code == 422 and "sedang melayani nomor lain" in r.json()["detail"]
    assert _ubah(k, ids[3], status="aktif", meja=2, dibuat=f"{t}T09:20:00", dipanggil=f"{t}T09:25:00").status_code == 200


def test_kelola_hapus(admin):
    from app import db
    k = admin
    t, ids = _seed_kemarin(k)
    lain = db.q("SELECT count(*) AS n FROM queue_tiket WHERE tanggal<>%s", (t,), one=True)["n"]
    assert k.post("/api/data/hapus", json={"ids": []}).status_code == 422
    assert k.post("/api/data/hapus", json={"ids": ["x"]}).status_code == 422
    assert k.post("/api/data/hapus", json={"ids": [999999999]}).json()["terhapus"] == 0
    r = k.post("/api/data/hapus", json={"ids": [ids[1], ids[3]]})
    assert r.status_code == 200 and r.json() == {"ok": True, "terhapus": 2}
    assert [x["nomor"] for x in k.get(f"/api/data?tanggal={t}").json()["rows"]] == [2]
    # soft delete: barisnya tetap ada di tabel, hanya diberi tanda
    sisa = db.q("SELECT nomor, dihapus_at IS NOT NULL AS hapus FROM queue_tiket WHERE tanggal=%s ORDER BY nomor", (t,))
    assert [(x["nomor"], x["hapus"]) for x in sisa] == [(1, True), (2, False), (3, True)]
    assert k.post("/api/data/hapus", json={"ids": [ids[1]]}).json()["terhapus"] == 0            # sudah terhapus: tidak dihitung lagi
    assert _ubah(k, ids[1], status="menunggu", dibuat=f"{t}T09:00:00").status_code == 404      # data terhapus tidak bisa diedit
    r = k.post("/api/data/hapus-tanggal", json={"tanggal": str(t)})
    assert r.json() == {"ok": True, "terhapus": 1} and k.get(f"/api/data?tanggal={t}").json()["rows"] == []
    assert db.q("SELECT count(*) AS n FROM queue_tiket WHERE tanggal=%s", (t,), one=True)["n"] == 3   # semua baris masih tersimpan
    assert db.q("SELECT count(*) AS n FROM queue_tiket WHERE tanggal=%s AND dihapus_at IS NOT NULL", (t,), one=True)["n"] == 3
    assert k.post("/api/data/hapus-tanggal", json={"tanggal": "bukan"}).status_code == 422
    assert db.q("SELECT count(*) AS n FROM queue_tiket WHERE tanggal<>%s AND dihapus_at IS NULL", (t,), one=True)["n"] == lain   # tanggal lain utuh
    assert k.srv.psql("select count(*) from queue_antrian_admisi").split()[2] == "1"                      # tabel lama utuh


def test_data_terhapus_hilang_dari_laporan(riwayat, mod):
    k = riwayat
    sebelum = hitung(mod, jenis="harian", tanggal=__import__("datetime").date(2026, 8, 3))["ring"]
    assert sebelum["total"] == 3
    k.srv.psql("UPDATE queue_tiket SET dihapus_at=now() WHERE tanggal='2026-08-03' AND nomor=3")      # 003 = tidak hadir
    h = hitung(mod, jenis="harian", tanggal=__import__("datetime").date(2026, 8, 3))["ring"]
    assert (h["total"], h["selesai"], h["lewat"]) == (2, 2, 0)
    b = hitung(mod, jenis="bulanan", bulan=8, tahun=2026)["ring"]
    assert (b["total"], b["lewat"]) == (4, 0)
    k.srv.psql("UPDATE queue_tiket SET dihapus_at=now() WHERE tanggal<'2026-09-01'")
    assert mod.laporan.opsi_tahun(mod.laporan.periode(jenis="tahunan", tahun=2026))[-1] == 2026       # tahun 2025 (terhapus) tidak lagi jadi pilihan
    assert hitung(mod, jenis="tahunan", tahun=2025)["ring"]["total"] == 0
    assert hitung(mod, jenis="tahunan", tahun=2026)["ring"]["total"] == 1
    k.srv.psql("UPDATE queue_tiket SET dihapus_at=NULL")


def test_migrasi_skema_lama_dan_nomor_unik_hanya_data_aktif(admin):
    from app import db
    k = admin
    # meniru tabel versi sebelumnya: tanpa kolom dihapus_at, dengan UNIQUE (tanggal, nomor)
    k.srv.psql("DROP INDEX IF EXISTS queue_tiket_nomor_aktif; ALTER TABLE queue_tiket DROP COLUMN dihapus_at;"
               "ALTER TABLE queue_tiket ADD CONSTRAINT queue_tiket_tanggal_nomor_key UNIQUE (tanggal, nomor)")
    for _ in range(2):                                                         # idempoten: aman dijalankan tiap start
        for s in db.TABEL:
            db.q(s)
    ada = k.srv.psql("select column_name from information_schema.columns where table_name='queue_tiket' and column_name='dihapus_at'")
    assert "dihapus_at" in ada
    assert "queue_tiket_tanggal_nomor_key" not in k.srv.psql("select conname from pg_constraint where conrelid='queue_tiket'::regclass")
    bersih(k)
    assert k.post("/api/tiket").json()["nomor"] == 1
    k.srv.psql("UPDATE queue_tiket SET dihapus_at=now()")
    assert k.post("/api/tiket").json()["nomor"] == 1                             # nomor bekas data terhapus boleh dipakai lagi
    assert db.q("SELECT count(*) AS n FROM queue_tiket WHERE nomor=1", one=True)["n"] == 2
    assert k.srv.psql("select count(*) from queue_tiket where tanggal=current_date and nomor=1 and dihapus_at is null").split()[2] == "1"
    bersih(k)


def test_kelola_hapus_hari_ini_nomor_mulai_dari_satu(admin):
    k = admin
    bersih(k)
    assert [k.post("/api/tiket").json()["nomor"] for _ in range(3)] == [1, 2, 3]
    hari = k.get("/api/data").json()
    assert len(hari["rows"]) == 3 and [r["status"] for r in hari["rows"]] == ["menunggu"] * 3
    assert k.post("/api/data/hapus-tanggal", json={"tanggal": hari["tanggal"]}).json()["terhapus"] == 3
    assert k.get("/api/loket/1").json()["jumlah"] == 0 and k.get("/api/monitor").json()["sisa"] == 0
    assert k.post("/api/tiket").json()["nomor"] == 1                    # penomoran hari ini mulai lagi dari 1
    bersih(k)
