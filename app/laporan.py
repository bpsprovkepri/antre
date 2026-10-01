"""Laporan lama layanan: harian, bulanan, tahunan, dan rentang tanggal (+ ekspor Excel)."""
import calendar
import io
from datetime import date, timedelta
from urllib.parse import urlencode

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import pengaturan
from .config import TZ
from .db import q
from .util import BULAN, dur, hari_ini, hms, mmss, p3, tanggal_id

JENIS = ("harian", "bulanan", "tahunan", "rentang")
MAKS_HARI = 366
MAKS_DETAIL_LAYAR = 300      # jumlah baris detail yang ditampilkan di layar (selengkapnya ada di Excel)
BULAN_PENDEK = ["", "Jan", "Feb", "Mar", "Apr", "Mei", "Jun", "Jul", "Agu", "Sep", "Okt", "Nov", "Des"]
HARI_PENDEK = ["Sen", "Sel", "Rab", "Kam", "Jum", "Sab", "Min"]


class PeriodeError(ValueError):
    pass


def _bulan_lalu(t, m, n):
    idx = t * 12 + (m - 1) + n
    return idx // 12, idx % 12 + 1


def periode(jenis="harian", tanggal=None, bulan=None, tahun=None, awal=None, akhir=None) -> dict:
    """Ubah pilihan pengguna menjadi rentang tanggal [awal, akhir]. Melempar PeriodeError bila tidak valid."""
    hi = hari_ini()
    jenis = jenis if jenis in JENIS else "harian"
    per = {"jenis": jenis}
    if jenis == "harian":
        d = tanggal or hi
        a = b = d
        per.update(judul="Laporan Harian", sub=tanggal_id(d), kode=f"harian-{d}", tanggal=d.isoformat())
        qs = {"jenis": jenis, "tanggal": d.isoformat()}
    elif jenis == "bulanan":
        t, m = tahun or hi.year, bulan or hi.month
        if not 1 <= m <= 12:
            raise PeriodeError("Bulan harus 1 sampai 12")
        if not 2000 <= t <= 2100:
            raise PeriodeError("Tahun harus antara 2000 dan 2100")
        a, b = date(t, m, 1), date(t, m, calendar.monthrange(t, m)[1])
        per.update(judul="Laporan Bulanan", sub=f"{BULAN[m]} {t}", kode=f"bulanan-{t}-{m:02d}", bulan=m, tahun=t)
        qs = {"jenis": jenis, "bulan": m, "tahun": t}
    elif jenis == "tahunan":
        t = tahun or hi.year
        if not 2000 <= t <= 2100:
            raise PeriodeError("Tahun harus antara 2000 dan 2100")
        a, b = date(t, 1, 1), date(t, 12, 31)
        per.update(judul="Laporan Tahunan", sub=f"Tahun {t}", kode=f"tahunan-{t}", tahun=t)
        qs = {"jenis": jenis, "tahun": t}
    else:
        a, b = awal or hi.replace(day=1), akhir or hi
        if b < a:
            raise PeriodeError("Tanggal akhir tidak boleh lebih awal dari tanggal awal")
        if (b - a).days + 1 > MAKS_HARI:
            raise PeriodeError(f"Rentang maksimal {MAKS_HARI} hari")
        if a.year < 2000 or b.year > 2100:
            raise PeriodeError("Tanggal di luar batas yang didukung")
        per.update(judul="Laporan Rentang Waktu", sub=f"{tanggal_id(a)} – {tanggal_id(b)}", kode=f"rentang-{a}_{b}")
        qs = {"jenis": jenis, "awal": a.isoformat(), "akhir": b.isoformat()}
    per.update(awal=a, akhir=b, hari=(b - a).days + 1, awal_iso=a.isoformat(), akhir_iso=b.isoformat(), qs=urlencode(qs))
    return per


def navigasi(per: dict) -> dict:
    """Tautan sebelumnya/berikutnya dan pintasan rentang."""
    hi = hari_ini()
    nav = {"prev": None, "next": None, "kini": None, "preset": []}
    j = per["jenis"]
    if j == "harian":
        d = per["awal"]
        nav["prev"] = urlencode({"jenis": j, "tanggal": (d - timedelta(days=1)).isoformat()})
        nav["next"] = urlencode({"jenis": j, "tanggal": (d + timedelta(days=1)).isoformat()}) if d < hi else None
        nav["kini"] = urlencode({"jenis": j, "tanggal": hi.isoformat()})
    elif j == "bulanan":
        t, m = per["tahun"], per["bulan"]
        pt, pm = _bulan_lalu(t, m, -1)
        nt, nm = _bulan_lalu(t, m, 1)
        nav["prev"] = urlencode({"jenis": j, "bulan": pm, "tahun": pt})
        nav["next"] = urlencode({"jenis": j, "bulan": nm, "tahun": nt}) if (t, m) < (hi.year, hi.month) else None
        nav["kini"] = urlencode({"jenis": j, "bulan": hi.month, "tahun": hi.year})
    elif j == "tahunan":
        t = per["tahun"]
        nav["prev"] = urlencode({"jenis": j, "tahun": t - 1})
        nav["next"] = urlencode({"jenis": j, "tahun": t + 1}) if t < hi.year else None
        nav["kini"] = urlencode({"jenis": j, "tahun": hi.year})
    else:
        def r(a, b):
            return urlencode({"jenis": "rentang", "awal": a.isoformat(), "akhir": b.isoformat()})
        nav["preset"] = [("7 hari terakhir", r(hi - timedelta(days=6), hi)), ("30 hari terakhir", r(hi - timedelta(days=29), hi)),
                         ("Bulan ini", r(hi.replace(day=1), hi)), ("Tahun ini", r(date(hi.year, 1, 1), hi))]
    return nav


def opsi_tahun(per: dict) -> list:
    r = q("SELECT min(tanggal) AS a FROM queue_tiket", one=True)
    awal = min(r["a"].year if r and r["a"] else hari_ini().year, per.get("tahun") or hari_ini().year)
    return list(range(hari_ini().year, awal - 1, -1))


# ---------- perhitungan ----------
def _rata(x):
    return sum(x) / len(x) if x else None


def _granularitas(per):
    if per["jenis"] == "harian":
        return "jam"
    if per["jenis"] == "tahunan" or per["hari"] > 62:
        return "bulan"
    return "hari"


def _label(gran, k, per):
    if gran == "jam":
        return f"{k:02d}", f"Pukul {k:02d}.00–{k:02d}.59"
    if gran == "hari":
        beda_bulan = per["awal"].month != per["akhir"].month or per["awal"].year != per["akhir"].year
        pendek = f"{k.day:02d}/{k.month:02d}" if beda_bulan else f"{k.day:02d}"
        return pendek, f"{HARI_PENDEK[k.weekday()]}, {k.day} {BULAN[k.month]} {k.year}"
    beda_tahun = per["awal"].year != per["akhir"].year
    return (f"{BULAN_PENDEK[k[1]]} {str(k[0])[2:]}" if beda_tahun else BULAN_PENDEK[k[1]]), f"{BULAN[k[1]]} {k[0]}"


def _kunci_semua(gran, per, jam_ada):
    if gran == "jam":
        return list(range(min(7, min(jam_ada)), max(16, max(jam_ada)) + 1)) if jam_ada else []
    if gran == "hari":
        return [per["awal"] + timedelta(days=i) for i in range(per["hari"])]
    hasil, (t, m) = [], (per["awal"].year, per["awal"].month)
    while (t, m) <= (per["akhir"].year, per["akhir"].month):
        hasil.append((t, m))
        t, m = _bulan_lalu(t, m, 1)
    return hasil


def data(per: dict, p: dict) -> dict:
    hi = hari_ini()
    rows = q("SELECT * FROM queue_tiket WHERE tanggal BETWEEN %s AND %s ORDER BY tanggal, nomor", (per["awal"], per["akhir"]))
    gran = _granularitas(per)
    lay, tgg, per_meja, hari_aktif, bucket, jam_ada = [], [], {}, set(), {}, set()
    for r in rows:
        r["tunggu"] = dur(r["dibuat_at"], r["dipanggil_at"])
        r["layanan"] = dur(r["mulai_at"], r["selesai_at"]) if (r["selesai_at"] and not r["lewat"]) else None
        lampau = r["tanggal"] < hi
        r["status"] = ("tidak dilayani" if lampau else "menunggu") if not r["dipanggil_at"] else \
                      ("belum ditutup" if lampau else "berjalan") if not r["selesai_at"] else \
                      "tidak hadir" if r["lewat"] else "selesai"
        hari_aktif.add(r["tanggal"])
        if r["tunggu"] is not None:
            tgg.append(r["tunggu"])
        if r["layanan"] is not None:
            lay.append(r["layanan"])
        if gran == "jam":
            k = r["dibuat_at"].astimezone(TZ).hour
            jam_ada.add(k)
        elif gran == "hari":
            k = r["tanggal"]
        else:
            k = (r["tanggal"].year, r["tanggal"].month)
        b = bucket.setdefault(k, {"n": 0, "lay": [], "tgg": [], "lewat": 0, "tdk": 0})
        b["n"] += 1
        b["lewat"] += 1 if r["lewat"] else 0
        b["tdk"] += 1 if r["status"] in ("tidak dilayani", "belum ditutup") else 0
        if r["layanan"] is not None:
            b["lay"].append(r["layanan"])
        if r["tunggu"] is not None:
            b["tgg"].append(r["tunggu"])
        if r["meja"]:
            m = per_meja.setdefault(r["meja"], {"n": 0, "lay": [], "lewat": 0})
            m["n"] += 1
            m["lewat"] += 1 if r["lewat"] else 0
            if r["layanan"] is not None:
                m["lay"].append(r["layanan"])

    ring = {"total": len(rows), "selesai": len(lay), "lewat": sum(1 for r in rows if r["lewat"]),
            "menunggu": sum(1 for r in rows if r["status"] == "menunggu"),
            "berjalan": sum(1 for r in rows if r["status"] == "berjalan"),
            "tdk_dilayani": sum(1 for r in rows if r["status"] in ("tidak dilayani", "belum ditutup")),
            "rata_layanan": mmss(_rata(lay)), "maks_layanan": mmss(max(lay)) if lay else "-",
            "min_layanan": mmss(min(lay)) if lay else "-", "rata_tunggu": mmss(_rata(tgg)),
            "maks_tunggu": mmss(max(tgg)) if tgg else "-",
            "hari_aktif": len(hari_aktif), "rata_per_hari": round(len(rows) / len(hari_aktif), 1) if hari_aktif else 0}

    seri = []
    for k in _kunci_semua(gran, per, jam_ada):
        b = bucket.get(k, {"n": 0, "lay": [], "tgg": [], "lewat": 0, "tdk": 0})
        pendek, panjang = _label(gran, k, per)
        ra = _rata(b["lay"])
        seri.append({"label": pendek, "label_panjang": panjang, "n": b["n"], "selesai": len(b["lay"]), "lewat": b["lewat"],
                     "tdk": b["tdk"], "rata": mmss(ra), "rata_detik": ra or 0, "maks": mmss(max(b["lay"])) if b["lay"] else "-",
                     "rata_tunggu": mmss(_rata(b["tgg"]))})
    mn, mr = max([s["n"] for s in seri] + [1]), max([s["rata_detik"] for s in seri] + [1])
    for s in seri:
        s["pct"], s["pct_lay"] = round(s["n"] * 100 / mn), round(s["rata_detik"] * 100 / mr)

    def rata(x):
        return mmss(_rata(x))
    meja = [{"nama": pengaturan.nama_meja(p, no) or f"Meja {no}", "n": v["n"], "selesai": len(v["lay"]), "lewat": v["lewat"],
             "rata": rata(v["lay"]), "maks": mmss(max(v["lay"])) if v["lay"] else "-"} for no, v in sorted(per_meja.items())]
    for r in rows:
        r.update(nomor_txt=p3(r["nomor"]), tgl_txt=r["tanggal"].strftime("%d/%m/%Y"), ambil=hms(r["dibuat_at"]),
                 panggil=hms(r["dipanggil_at"]), selesai_jam=hms(r["selesai_at"]), tunggu_txt=mmss(r["tunggu"]),
                 layanan_txt=mmss(r["layanan"]),
                 meja_nama=(pengaturan.nama_meja(p, r["meja"]) or f"Meja {r['meja']}") if r["meja"] else "-")
    return {"ring": ring, "seri": seri, "gran": gran, "meja": meja, "rows": rows,
            "tampil_detail": len(rows) <= MAKS_DETAIL_LAYAR,
            "judul_seri": {"jam": "per jam", "hari": "per hari", "bulan": "per bulan"}[gran]}


# ---------- Excel ----------
def _gaya(ws, lebar):
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="0A6B38")
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for i, w in enumerate(lebar, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"


def xlsx(per: dict, p: dict) -> bytes:
    d = data(per, p)
    g = d["ring"]
    wb = Workbook()
    ws = wb.active
    ws.title = "Ringkasan"
    ws.append(["Keterangan", "Nilai"])
    for baris in ([per["judul"], per["sub"]], ["Instansi", p["nama_instansi"]], ["Total antrian", g["total"]],
                  ["Selesai dilayani", g["selesai"]], ["Tidak hadir", g["lewat"]],
                  ["Tidak dilayani / belum ditutup", g["tdk_dilayani"]],
                  ["Rata-rata lama layanan (mm:ss)", g["rata_layanan"]], ["Layanan terlama (mm:ss)", g["maks_layanan"]],
                  ["Layanan tercepat (mm:ss)", g["min_layanan"]], ["Rata-rata waktu tunggu (mm:ss)", g["rata_tunggu"]],
                  ["Hari beroperasi", g["hari_aktif"]], ["Rata-rata antrian per hari", g["rata_per_hari"]], [],
                  ["Meja", "Dilayani", "Tidak hadir", "Rata-rata layanan", "Terlama"]):
        ws.append(baris)
    for m in d["meja"]:
        ws.append([m["nama"], m["selesai"], m["lewat"], m["rata"], m["maks"]])
    _gaya(ws, [36, 26, 16, 20, 14])
    r2 = wb.create_sheet("Rincian")
    r2.append([{"jam": "Jam", "hari": "Tanggal", "bulan": "Bulan"}[d["gran"]], "Total antrian", "Selesai dilayani", "Tidak hadir",
               "Tidak dilayani / belum ditutup", "Rata-rata layanan (mm:ss)", "Layanan terlama (mm:ss)", "Rata-rata layanan (detik)",
               "Rata-rata tunggu (mm:ss)"])
    for s in d["seri"]:
        r2.append([s["label_panjang"], s["n"], s["selesai"], s["lewat"], s["tdk"], s["rata"], s["maks"],
                   round(s["rata_detik"]) if s["selesai"] else None, s["rata_tunggu"]])
    _gaya(r2, [30, 14, 16, 12, 22, 22, 22, 22, 22])
    r3 = wb.create_sheet("Detail")
    r3.append(["Tanggal", "Nomor", "Meja", "Ambil", "Panggil & Mulai", "Selesai", "Tunggu (mm:ss)", "Layanan (mm:ss)",
               "Tunggu (detik)", "Layanan (detik)", "Status"])
    for r in d["rows"]:
        r3.append([r["tanggal"], r["nomor_txt"], r["meja_nama"], r["ambil"], r["panggil"], r["selesai_jam"], r["tunggu_txt"],
                   r["layanan_txt"], r["tunggu"], r["layanan"], r["status"]])
    for row in r3.iter_rows(min_row=2, max_col=1):
        row[0].number_format = "DD/MM/YYYY"
    _gaya(r3, [14, 10, 16, 11, 16, 11, 16, 16, 14, 14, 16])
    r3.auto_filter.ref = r3.dimensions
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
